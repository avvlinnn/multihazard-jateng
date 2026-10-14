"""
Baca semua sheet (2016-2025) di Data_Banjir-Longsor_Jateng_2016-2025.xlsx,
filter cuma kejadian BANJIR & TANAH LONGSOR, satukan jadi satu tabel bersih.

Ada 2 format beda di file ini:
  - Format "pusdalops" (sheet 2016-2023): kolom jenis bencana one-hot
    (kolom "BANJIR" isinya 1 kalau kejadian, kolom "TANAH LONGSOR" isinya 1
    kalau longsor). Kolom desa/kecamatan namanya "Ds." dan "KECAMATAN".
  - Format "modern" (sheet 2024-2025): kolom jenis bencana jadi satu teks
    di kolom "Jenis Bencana" (misal "101.Banjir" / "102.Tanah Longsor").
    Kolom desa/kecamatan namanya "Desa" dan "Kecamatan".

Script ini otomatis deteksi baris header & format tiap sheet (bukan hardcode
nomor baris), soalnya posisi header beda-beda dikit tiap tahun.

Output:
  - bpbd_banjir_longsor_bersih.csv   -> data yang tanggalnya berhasil dibaca
  - bpbd_gagal_parse_tanggal.csv     -> baris yang tanggalnya GAGAL dibaca,
                                        perlu dicek & dibenerin manual satu-satu
"""

import re
import datetime as dt
import openpyxl
import pandas as pd

# ==== KONFIGURASI ====
INPUT_FILE = "data_mentah/pusdalops/Data Banjir-Longsor Jateng 2016-2025.xlsx"
OUTPUT_BERSIH = "data_olahan/bpbd_banjir_longsor_bersih.csv"
OUTPUT_GAGAL_TANGGAL = "data_olahan/bpbd_gagal_parse_tanggal.csv"
# ======================

BULAN_ID = {
    "januari": 1, "februari": 2, "maret": 3, "april": 4, "mei": 5, "juni": 6,
    "juli": 7, "agustus": 8, "september": 9, "oktober": 10,
    "november": 11, "desember": 12,
}


def norm(x):
    """Bikin nama kolom jadi lowercase & rapi, buat matching antar sheet."""
    return str(x).strip().lower() if x is not None else ""


def cari_baris_header(ws, max_cek=12):
    """
    Cari baris header dengan nyari baris yang punya kolom 'TGL.KEJADIAN'
    (format pusdalops) atau 'Tarikh' (format modern).
    Return: (index_baris_header_1based, format_asal, header_row_values)
    """
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=max_cek, values_only=True), start=1):
        normed = [norm(c) for c in row]
        if "tgl.kejadian" in normed:
            return i, "pusdalops", row
        if "tarikh" in normed:
            return i, "modern", row
    return None, None, None


def cari_kolom(header_row, target_norm):
    """Cari index kolom yang nama (setelah di-norm) match persis target_norm."""
    for idx, val in enumerate(header_row):
        if norm(val) == target_norm:
            return idx
    return None


def cari_kolom_mengandung(header_row, kata_kata):
    """Cari index kolom pertama yang nama-nya mengandung semua kata di kata_kata."""
    for idx, val in enumerate(header_row):
        n = norm(val)
        if all(k in n for k in kata_kata):
            return idx
    return None


def parse_tanggal(value, tahun_sheet):
    """Coba parse macam-macam bentuk tanggal. Return datetime.date atau None kalau gagal."""
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, (int, float)):
        try:
            return dt.date(1899, 12, 30) + dt.timedelta(days=int(value))
        except (ValueError, OverflowError):
            return None
    if isinstance(value, str):
        s = value.strip()
        s = re.sub(r"(pkl|pukul).*", "", s, flags=re.IGNORECASE).strip()

        # pola dd/mm/yyyy atau d-m-yy (boleh ada spasi di sekitar / atau -)
        m = re.search(r"(\d{1,2})\s*[/-]\s*(\d{1,2})\s*[/-]\s*(\d{2,4})", s)
        if m:
            try:
                tgl_bersih = f"{m.group(1)}/{m.group(2)}/{m.group(3)}"
                return pd.to_datetime(tgl_bersih, dayfirst=True).date()
            except (ValueError, TypeError):
                return None

        # pola "1 Januari 2023"
        m = re.search(r"(\d{1,2})\s+([A-Za-z]+)\s+(\d{4})", s)
        if m:
            bulan = BULAN_ID.get(m.group(2).strip().lower())
            if bulan:
                try:
                    return dt.date(int(m.group(3)), bulan, int(m.group(1)))
                except ValueError:
                    return None

        return None
    return None


def proses_sheet_pusdalops(ws, header_idx, header_row, tahun_sheet):
    idx_tgl = cari_kolom(header_row, "tgl.kejadian")
    idx_kabkota = cari_kolom(header_row, "kab/kota")
    idx_banjir = cari_kolom(header_row, "banjir")
    idx_longsor = cari_kolom(header_row, "tanah longsor")
    idx_desa = cari_kolom(header_row, "ds.")
    idx_kecamatan = cari_kolom(header_row, "kecamatan")
    idx_keterangan = cari_kolom(header_row, "keterangan")

    hasil = []
    for row in ws.iter_rows(min_row=header_idx + 1, values_only=True):
        if idx_banjir is None or idx_longsor is None:
            break
        is_banjir = row[idx_banjir] not in (None, "", 0) if idx_banjir is not None else False
        is_longsor = row[idx_longsor] not in (None, "", 0) if idx_longsor is not None else False
        if not (is_banjir or is_longsor):
            continue

        if is_banjir and is_longsor:
            jenis = "Banjir dan Longsor"
        elif is_banjir:
            jenis = "Banjir"
        else:
            jenis = "Longsor"

        hasil.append({
            "tanggal_raw": row[idx_tgl] if idx_tgl is not None else None,
            "tahun_sheet": tahun_sheet,
            "format_asal": "pusdalops",
            "kab_kota": row[idx_kabkota] if idx_kabkota is not None else None,
            "kecamatan": row[idx_kecamatan] if idx_kecamatan is not None else None,
            "desa": row[idx_desa] if idx_desa is not None else None,
            "jenis_bencana": jenis,
            "keterangan": row[idx_keterangan] if idx_keterangan is not None else None,
        })
    return hasil


def proses_sheet_modern(ws, header_idx, header_row, tahun_sheet):
    idx_tgl = cari_kolom(header_row, "tarikh")
    idx_kabkota = cari_kolom_mengandung(header_row, ["kab", "kota"])
    idx_kecamatan = cari_kolom(header_row, "kecamatan")
    idx_desa = cari_kolom(header_row, "desa")
    idx_jenis = cari_kolom(header_row, "jenis bencana")
    idx_keterangan = cari_kolom(header_row, "kronologis")
    if idx_keterangan is None:
        idx_keterangan = cari_kolom(header_row, "deskripsi")

    hasil = []
    for row in ws.iter_rows(min_row=header_idx + 1, values_only=True):
        if idx_jenis is None:
            break
        teks_jenis = norm(row[idx_jenis]) if idx_jenis < len(row) else ""
        is_banjir = "banjir" in teks_jenis
        is_longsor = "longsor" in teks_jenis
        if not (is_banjir or is_longsor):
            continue

        if is_banjir and is_longsor:
            jenis = "Banjir dan Longsor"
        elif is_banjir:
            jenis = "Banjir"
        else:
            jenis = "Longsor"

        hasil.append({
            "tanggal_raw": row[idx_tgl] if idx_tgl is not None else None,
            "tahun_sheet": tahun_sheet,
            "format_asal": "modern",
            "kab_kota": row[idx_kabkota] if idx_kabkota is not None else None,
            "kecamatan": row[idx_kecamatan] if idx_kecamatan is not None else None,
            "desa": row[idx_desa] if idx_desa is not None else None,
            "jenis_bencana": jenis,
            "keterangan": row[idx_keterangan] if idx_keterangan is not None else None,
        })
    return hasil


def main():
    wb = openpyxl.load_workbook(INPUT_FILE, read_only=True, data_only=True)
    semua_baris = []

    for nama_sheet in wb.sheetnames:
        ws = wb[nama_sheet]
        header_idx, fmt, header_row = cari_baris_header(ws)
        if header_idx is None:
            print(f"⚠️  Sheet {nama_sheet}: gagal nemu baris header, dilewati.")
            continue

        if fmt == "pusdalops":
            baris = proses_sheet_pusdalops(ws, header_idx, header_row, nama_sheet)
        else:
            baris = proses_sheet_modern(ws, header_idx, header_row, nama_sheet)

        print(f"Sheet {nama_sheet} ({fmt}): {len(baris)} baris banjir/longsor")
        semua_baris.extend(baris)

    df = pd.DataFrame(semua_baris)

    # parse tanggal
    df["tanggal"] = df.apply(lambda r: parse_tanggal(r["tanggal_raw"], r["tahun_sheet"]), axis=1)

    gagal = df[df["tanggal"].isna()].copy()
    bersih = df[df["tanggal"].notna()].copy()

    kolom_akhir = ["tanggal", "tahun_sheet", "format_asal", "kab_kota",
                   "kecamatan", "desa", "jenis_bencana", "keterangan"]
    bersih = bersih[kolom_akhir]
    gagal = gagal[["tanggal_raw"] + kolom_akhir[1:]]

    bersih.to_csv(OUTPUT_BERSIH, index=False)
    gagal.to_csv(OUTPUT_GAGAL_TANGGAL, index=False)

    print(f"\nTotal baris banjir/longsor: {len(df)}")
    print(f"  - berhasil parse tanggal -> {OUTPUT_BERSIH} ({len(bersih)} baris)")
    print(f"  - GAGAL parse tanggal, perlu dicek manual -> {OUTPUT_GAGAL_TANGGAL} ({len(gagal)} baris)")
    print("\nRingkasan jenis bencana:")
    print(bersih["jenis_bencana"].value_counts())
    print("\nCatatan penting:")
    print("- Kolom 'kecamatan' & 'desa' MASIH mentah (belum dipecah kalau isinya lebih dari satu nilai).")
    print("  Lanjut jalanin split_multi_desa.py setelah ini.")
    print("- Nama kab/kota MASIH mentah (belum disamain ejaannya / dicek duplikat Kab vs Kota).")
    print("- Kolom 'TANAH GERAK'/'Tanah Gerak' di sheet pusdalops TIDAK ikut dihitung sebagai longsor,")
    print("  cuma kolom 'TANAH LONGSOR' aja. Bilang kalau mau itu diikutin juga.")


if __name__ == "__main__":
    main()